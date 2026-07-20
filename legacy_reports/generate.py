# -*- coding: utf-8 -*-
"""
Al-Qavi — Legacy (desktop Trade) report generator.
Reads TDistribution.mdb (MS Access) directly and reproduces the core
distribution/accounting reports from the double-entry ledger (VoucherDetail):
  Trial Balance, Recovery (Receivables), Income Statement, Balance Sheet,
  Sales Register, and a full General Ledger.

Output: an Excel workbook + printable HTML (Income Statement, Recovery).
No dependency on the desktop app; read-only.
"""
import os, re, datetime
from collections import defaultdict
from access_parser import AccessParser
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DB_PATH = r'D:\al qavi glt\TDistribution.mdb'
OUT_DIR = r'D:\al qavi glt\WebReports'
os.makedirs(OUT_DIR, exist_ok=True)

db = AccessParser(DB_PATH)

def tbl(name):
    d = db.parse_table(name)
    cols = list(d.keys())
    n = len(d[cols[0]]) if cols else 0
    return [{c: d[c][i] for c in cols} for i in range(n)]

def f(x):
    try:
        return float(str(x).replace(',', '').strip())
    except Exception:
        return 0.0

def pdate(v):
    try:
        v = int(v)
        if v <= 0:
            return ''
        return '%02d-%02d-%04d' % (v % 100, (v // 100) % 100, v // 10000)
    except Exception:
        return ''

# ── Load master + ledger ────────────────────────────────────────────────
main_acc = {int(r['MainAccID']): (r['MainAccName'] or '').strip() for r in tbl('MainAccount')}
sub_acc  = {int(r['SubAccID']): ((r['SubAccName'] or '').strip(), int(r['MainAccID'])) for r in tbl('Accounts2L')}
l3_acc   = {int(r['Acc3LID']): ((r['Acc3LName'] or '').strip(), int(r['SubAccID'])) for r in tbl('Accounts3L')}
areas    = {int(r['AreaID']): (r['AreaName'] or '').strip() for r in tbl('Area')}
staff    = {int(r['StaffID']): (r['StaffName'] or '').strip() for r in tbl('Staff')}
accounts = {}
for r in tbl('Accounts'):
    accounts[int(r['AccID'])] = {
        'name': (r['AccName'] or '').strip(),
        'area': areas.get(int(r['AreaID'] or 0), ''),
        'l3': int(str(r['Acc3LID'] or 0) or 0),
        'contact': (r['ContactPerson'] or '').strip(),
        'cell': (r['CellNo'] or '').strip(),
    }

def classify(accid):
    """Return (main_id, main_name, sub_name, l3_name) for a leaf AccID."""
    l3id = int(accid) // 10000
    l3name, subid = l3_acc.get(l3id, ('', 0))
    subname, mainid = sub_acc.get(subid, ('', 0))
    return mainid, main_acc.get(mainid, ''), subname, l3name

def acc_name(accid):
    a = accounts.get(int(accid))
    if a and a['name']:
        return a['name']
    _, _, _, l3 = classify(accid)
    return l3 or ('Acc %d' % int(accid))

vd = tbl('VoucherDetail')
for r in vd:
    r['_acc'] = int(r['AccID']); r['_d'] = f(r['Debit']); r['_c'] = f(r['Credit'])
    r['_dt'] = int(r['DateVoch'] or 0)

# Per-account balances
bal = defaultdict(lambda: [0.0, 0.0])   # accid -> [debit, credit]
for r in vd:
    bal[r['_acc']][0] += r['_d']; bal[r['_acc']][1] += r['_c']

# ── Styling helpers ─────────────────────────────────────────────────────
H_FILL = PatternFill('solid', fgColor='232F3E')
H_FONT = Font(bold=True, color='FFFFFF', size=10)
T_FONT = Font(bold=True, size=10)
MONEY = '#,##0.00'
thin = Side(style='thin', color='D5D9D9')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def sheet(wb, title, headers, rows, money_cols=(), total_row=None, subtitle=None):
    ws = wb.create_sheet(title[:31])
    r0 = 1
    ws.cell(1, 1, 'AL-QAVI TRADERS').font = Font(bold=True, size=14, color='232F3E')
    ws.cell(2, 1, title).font = Font(bold=True, size=11)
    if subtitle:
        ws.cell(3, 1, subtitle).font = Font(size=9, color='777777')
        r0 = 5
    else:
        r0 = 4
    for j, h in enumerate(headers, 1):
        c = ws.cell(r0, j, h); c.fill = H_FILL; c.font = H_FONT; c.border = BORDER
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(r0 + i, j, v); c.border = BORDER
            if j in money_cols:
                c.number_format = MONEY; c.alignment = Alignment(horizontal='right')
    if total_row:
        rr = r0 + len(rows) + 1
        for j, v in enumerate(total_row, 1):
            c = ws.cell(rr, j, v); c.font = T_FONT
            c.fill = PatternFill('solid', fgColor='F0F2F2'); c.border = BORDER
            if j in money_cols:
                c.number_format = MONEY; c.alignment = Alignment(horizontal='right')
    # widths
    for j, h in enumerate(headers, 1):
        w = max(12, min(46, len(str(h)) + 2, max([len(str(r[j-1])) for r in rows[:400]] + [len(str(h))]) + 2))
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    return ws

# ── Build workbook ──────────────────────────────────────────────────────
wb = openpyxl.Workbook()
wb.remove(wb.active)
stamp = datetime.datetime.fromtimestamp(os.path.getmtime(DB_PATH)).strftime('%d-%m-%Y')
sub = 'Data source: desktop Trade DB · as of %s' % stamp

# 1) TRIAL BALANCE ------------------------------------------------------
tb_rows = []
tot_d = tot_c = 0.0
for accid in sorted(bal):
    d, c = bal[accid]
    net = d - c
    if abs(net) < 0.005 and d == 0 and c == 0:
        continue
    mid, mname, sname, l3 = classify(accid)
    side = 'Dr' if net >= 0 else 'Cr'
    tb_rows.append([accid, acc_name(accid), l3, mname, round(d, 2), round(c, 2), round(abs(net), 2), side])
    tot_d += d; tot_c += c
sheet(wb, 'Trial Balance',
      ['Code', 'Account', 'Head', 'Class', 'Debit', 'Credit', 'Balance', 'Dr/Cr'],
      tb_rows, money_cols=(5, 6, 7),
      total_row=['', 'TOTAL', '', '', round(tot_d, 2), round(tot_c, 2), '', ''], subtitle=sub)

# 2) RECOVERY (Receivables) --------------------------------------------
# Customers = accounts whose 3L head is Accounts Receivables; balance Dr = they owe us.
AR_HEADS = {k for k, (nm, _) in l3_acc.items() if 'receivable' in nm.lower()}
rec_rows = []
rec_total = 0.0
for accid, a in accounts.items():
    if a['l3'] in AR_HEADS:
        d, c = bal.get(accid, [0, 0]); net = d - c
        if net > 0.5:
            rec_rows.append([a['area'] or '—', accid, a['name'], a['contact'] or a['cell'], round(net, 2)])
            rec_total += net
rec_rows.sort(key=lambda r: (-r[4]))
sheet(wb, 'Recovery (Receivables)',
      ['Area', 'Code', 'Customer', 'Contact', 'Balance Due'],
      rec_rows, money_cols=(5,),
      total_row=['', '', 'TOTAL OUTSTANDING', '', round(rec_total, 2)], subtitle=sub)

# 3) INCOME STATEMENT --------------------------------------------------
rev = defaultdict(float); exp = defaultdict(float)
for accid, (d, c) in bal.items():
    mid, mname, sname, l3 = classify(accid)
    if mid == 4:      # Revenue: credit positive
        rev[l3 or acc_name(accid)] += (c - d)
    elif mid == 5:    # Expenses: debit positive
        exp[l3 or acc_name(accid)] += (d - c)
inc_rows = [['REVENUE', '']]
rev_tot = sum(rev.values())
for k in sorted(rev, key=lambda k: -rev[k]):
    inc_rows.append([k, round(rev[k], 2)])
inc_rows.append(['Total Revenue', round(rev_tot, 2)])
inc_rows.append(['', ''])
inc_rows.append(['EXPENSES', ''])
exp_tot = sum(exp.values())
for k in sorted(exp, key=lambda k: -exp[k]):
    inc_rows.append([k, round(exp[k], 2)])
inc_rows.append(['Total Expenses', round(exp_tot, 2)])
sheet(wb, 'Income Statement', ['Head', 'Amount'], inc_rows, money_cols=(2,),
      total_row=['NET PROFIT / (LOSS)', round(rev_tot - exp_tot, 2)], subtitle=sub)

# 4) BALANCE SHEET -----------------------------------------------------
grp = {1: defaultdict(float), 2: defaultdict(float), 3: defaultdict(float)}
for accid, (d, c) in bal.items():
    mid, mname, sname, l3 = classify(accid)
    if mid in grp:
        grp[mid][l3 or acc_name(accid)] += (d - c) if mid == 1 else (c - d)
bs_rows = []; assets = liab = eq = 0.0
for mid, label in [(1, 'ASSETS'), (2, 'LIABILITIES'), (3, 'EQUITY')]:
    bs_rows.append([label, ''])
    st = 0.0
    for k in sorted(grp[mid], key=lambda k: -grp[mid][k]):
        bs_rows.append([k, round(grp[mid][k], 2)]); st += grp[mid][k]
    bs_rows.append(['Total ' + label.title(), round(st, 2)]); bs_rows.append(['', ''])
    if mid == 1: assets = st
    elif mid == 2: liab = st
    else: eq = st
np_ = rev_tot - exp_tot
sheet(wb, 'Balance Sheet', ['Head', 'Amount'], bs_rows, money_cols=(2,),
      total_row=['Liab + Equity + Profit', round(liab + eq + np_, 2)], subtitle=sub + '  |  Assets: %.0f' % assets)

# 5) SALES REGISTER ----------------------------------------------------
sa = tbl('SaleAmount')
sa.sort(key=lambda r: (int(r['SaleDate'] or 0), str(r['SaleID'])))
sr_rows = []; s_amt = s_paid = 0.0
for r in sa:
    amt = f(r['Amount']); disp = f(r['DisPAmt']); disa = f(r['DiscAAmt']); paid = f(r['PaidCash'])
    net = amt - disp - disa
    sr_rows.append([str(r['SaleID']), pdate(r['SaleDate']), acc_name(r['AccID']),
                    staff.get(int(r['StaffID'] or 0), ''), round(amt, 2), round(disp + disa, 2),
                    round(net, 2), round(paid, 2)])
    s_amt += net; s_paid += paid
sheet(wb, 'Sales Register',
      ['Invoice', 'Date', 'Customer', 'Salesman', 'Gross', 'Discount', 'Net', 'Cash Recd'],
      sr_rows, money_cols=(5, 6, 7, 8),
      total_row=['', '', 'TOTAL', '', '', '', round(s_amt, 2), round(s_paid, 2)], subtitle=sub)

# 6) GENERAL LEDGER (all lines) ----------------------------------------
vt = {int(r['VTID']): (r['VTName'] or '').strip() for r in tbl('VoucherType')}
gl = sorted(vd, key=lambda r: (r['_acc'], r['_dt'], str(r['VoucherID']), int(r['LineID'] or 0)))
gl_rows = []; run = {}; cur = None
for r in gl:
    a = r['_acc']
    if a != cur:
        cur = a; run[a] = 0.0
    run[a] += r['_d'] - r['_c']
    gl_rows.append([a, acc_name(a), pdate(r['_dt']), str(r['VoucherID']),
                    vt.get(int(r['VT'] or 0), ''), (r['VDetail'] or '').strip(),
                    round(r['_d'], 2), round(r['_c'], 2), round(run[a], 2)])
sheet(wb, 'General Ledger',
      ['Code', 'Account', 'Date', 'Voucher', 'Type', 'Detail', 'Debit', 'Credit', 'Running Bal'],
      gl_rows, money_cols=(7, 8, 9), subtitle=sub)

# 7) PURCHASE REGISTER -------------------------------------------------
pa = tbl('PurchaseAmt')
pa.sort(key=lambda r: (int(r['PurDate'] or 0), str(r['PurID'])))
pr_rows = []; p_net = p_paid = 0.0
for r in pa:
    amt = f(r['PurAmt']); disc = f(r['DiscAmt']); stax = f(r['STaxAmt']); frt = f(r['Freight'])
    net = f(r['PurToAmt']) or (amt - disc + stax + frt); paid = f(r['PaidC'])
    pr_rows.append([str(r.get('BillNo') or ''), pdate(r['PurDate']), acc_name(r['AccID']),
                    round(amt, 2), round(disc, 2), round(stax, 2), round(frt, 2), round(net, 2), round(paid, 2)])
    p_net += net; p_paid += paid
sheet(wb, 'Purchase Register',
      ['Bill No', 'Date', 'Supplier', 'Amount', 'Discount', 'S.Tax', 'Freight', 'Net', 'Paid'],
      pr_rows, money_cols=(4, 5, 6, 7, 8, 9),
      total_row=['', '', 'TOTAL', '', '', '', '', round(p_net, 2), round(p_paid, 2)], subtitle=sub)

# 8) STOCK POSITION ----------------------------------------------------
comp = {int(r['CompID']): (r['CompName'] or '').strip() for r in tbl('Company')}
prod = {}
for r in tbl('Product'):
    prod[int(r['PID'])] = {'name': (r['ProdName'] or '').strip(),
                           'comp': comp.get(int(r['CompID'] or 0), ''), 'pack': (r['Packing'] or '')}
stock = defaultdict(lambda: [0.0, 0.0, 0.0])
for r in tbl('CompBatchStock'):
    pid = int(r['PID']); q = f(r['Qty'])
    stock[pid][0] += q; stock[pid][1] += q * f(r['UPRate']); stock[pid][2] += q * f(r['USRate'])
st_rows = []; st_cost = st_sale = 0.0
for pid, (q, cv, sv) in stock.items():
    if abs(q) < 0.001:
        continue
    p = prod.get(pid, {'name': 'PID %d' % pid, 'comp': '', 'pack': ''})
    st_rows.append([p['name'], p['comp'], p['pack'], round(q, 2), round(cv, 2), round(sv, 2)])
    st_cost += cv; st_sale += sv
st_rows.sort(key=lambda r: -r[4])
sheet(wb, 'Stock Position',
      ['Product', 'Company', 'Packing', 'Qty', 'Cost Value', 'Sale Value'],
      st_rows, money_cols=(5, 6),
      total_row=['', '', 'TOTAL', '', round(st_cost, 2), round(st_sale, 2)], subtitle=sub)

# 9) SALES ANALYSIS (salesman / area / month) --------------------------
by_sm = defaultdict(float); by_ar = defaultdict(float); by_mn = defaultdict(float)
for r in sa:
    net = f(r['Amount']) - f(r['DisPAmt']) - f(r['DiscAAmt'])
    by_sm[staff.get(int(r['StaffID'] or 0), '—')] += net
    by_ar[accounts.get(int(r['AccID']), {}).get('area', '—') or '—'] += net
    d = int(r['SaleDate'] or 0)
    by_mn['%04d-%02d' % (d // 10000, (d // 100) % 100)] += net
sheet(wb, 'Sales by Salesman', ['Salesman', 'Net Sales'],
      sorted([[k, round(v, 2)] for k, v in by_sm.items()], key=lambda x: -x[1]),
      money_cols=(2,), total_row=['TOTAL', round(sum(by_sm.values()), 2)], subtitle=sub)
sheet(wb, 'Sales by Area', ['Area', 'Net Sales'],
      sorted([[k, round(v, 2)] for k, v in by_ar.items()], key=lambda x: -x[1]),
      money_cols=(2,), total_row=['TOTAL', round(sum(by_ar.values()), 2)], subtitle=sub)
sheet(wb, 'Monthly Sales', ['Month', 'Net Sales'],
      sorted([[k, round(v, 2)] for k, v in by_mn.items()]),
      money_cols=(2,), total_row=['TOTAL', round(sum(by_mn.values()), 2)], subtitle=sub)

# 10) SALES BY PRODUCT (from line items) -------------------------------
by_pr = defaultdict(lambda: [0.0, 0.0])   # pid -> [qty, value]
for r in tbl('Sale'):
    pid = int(r['PID']); q = f(r['Qty']); val = q * f(r['USRate'])
    by_pr[pid][0] += q; by_pr[pid][1] += val
pr2 = []
for pid, (q, val) in by_pr.items():
    p = prod.get(pid, {'name': 'PID %d' % pid, 'comp': ''})
    pr2.append([p['name'], p['comp'], round(q, 2), round(val, 2)])
pr2.sort(key=lambda r: -r[3])
sheet(wb, 'Sales by Product', ['Product', 'Company', 'Qty Sold', 'Sale Value'],
      pr2, money_cols=(4,), total_row=['', 'TOTAL', '', round(sum(r[3] for r in pr2), 2)], subtitle=sub)

# order sheets nicely
wb.move_sheet('Trial Balance', -wb.sheetnames.index('Trial Balance'))
xlsx = os.path.join(OUT_DIR, 'AlQavi_Reports.xlsx')
wb.save(xlsx)

# ── Customer Ledgers workbook (per receivable party, running balance) ──
wb2 = openpyxl.Workbook(); wb2.remove(wb2.active)
led_ix = wb2.create_sheet('Index')
led_ix.cell(1, 1, 'AL-QAVI TRADERS — Customer Ledgers').font = Font(bold=True, size=13, color='232F3E')
led_ix.cell(2, 1, sub).font = Font(size=9, color='777777')
for j, h in enumerate(['Customer', 'Area', 'Closing Balance'], 1):
    c = led_ix.cell(4, j, h); c.fill = H_FILL; c.font = H_FONT
by_acc_lines = defaultdict(list)
for r in vd:
    by_acc_lines[r['_acc']].append(r)
ar_parties = sorted([aid for aid, a in accounts.items() if a['l3'] in AR_HEADS and (bal.get(aid, [0, 0])[0] - bal.get(aid, [0, 0])[1]) > 0.5],
                    key=lambda aid: -(bal[aid][0] - bal[aid][1]))
ix = 5
used = set()
for aid in ar_parties:
    a = accounts[aid]; closing = bal[aid][0] - bal[aid][1]
    nm = re.sub(r'[\\/*?:\[\]]', ' ', (a['name'] or ('Acc %d' % aid))).strip()[:28] or ('Acc %d' % aid)
    base = nm; k = 1
    while base in used:
        k += 1; base = ('%s %d' % (nm[:26], k))
    used.add(base)
    ws = wb2.create_sheet(base)
    ws.cell(1, 1, a['name']).font = Font(bold=True, size=12)
    ws.cell(2, 1, 'Area: %s   Contact: %s' % (a['area'], a['contact'] or a['cell'])).font = Font(size=9, color='555555')
    for j, h in enumerate(['Date', 'Voucher', 'Detail', 'Debit', 'Credit', 'Balance'], 1):
        c = ws.cell(4, j, h); c.fill = H_FILL; c.font = H_FONT; c.border = BORDER
    lines = sorted(by_acc_lines[aid], key=lambda r: (r['_dt'], str(r['VoucherID']), int(r['LineID'] or 0)))
    run = 0.0
    for i, r in enumerate(lines, 1):
        run += r['_d'] - r['_c']
        for j, v in enumerate([pdate(r['_dt']), str(r['VoucherID']), (r['VDetail'] or '').strip(),
                               round(r['_d'], 2), round(r['_c'], 2), round(run, 2)], 1):
            c = ws.cell(4 + i, j, v); c.border = BORDER
            if j in (4, 5, 6):
                c.number_format = MONEY; c.alignment = Alignment(horizontal='right')
    ws.column_dimensions['A'].width = 12; ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 40
    for col in 'DEF':
        ws.column_dimensions[col].width = 15
    ws.freeze_panes = 'A5'
    led_ix.cell(ix, 1, a['name']); led_ix.cell(ix, 2, a['area'])
    cc = led_ix.cell(ix, 3, round(closing, 2)); cc.number_format = MONEY
    ix += 1
for col, w in [('A', 32), ('B', 18), ('C', 16)]:
    led_ix.column_dimensions[col].width = w
xlsx2 = os.path.join(OUT_DIR, 'AlQavi_Customer_Ledgers.xlsx')
wb2.save(xlsx2)

# ── Printable HTML (Income Statement + Recovery) ───────────────────────
def money(x): return '{:,.2f}'.format(x)
css = """<style>body{font-family:Segoe UI,Arial,sans-serif;color:#111;margin:24px}
h1{font-size:20px;margin:0}h2{font-size:14px;color:#555;margin:2px 0 14px;font-weight:600}
table{border-collapse:collapse;width:100%;font-size:12.5px}th,td{border:1px solid #d5d9d9;padding:6px 9px}
th{background:#232F3E;color:#fff;text-align:left}td.r,th.r{text-align:right;font-variant-numeric:tabular-nums}
tr.sec td{background:#f0f2f2;font-weight:800;text-transform:uppercase;font-size:11px}
tr.tot td{font-weight:800;border-top:2px solid #232F3E}@media print{body{margin:0}}</style>"""

inc_html = ''.join(
    ('<tr class="sec"><td>%s</td><td class="r"></td></tr>' % r[0]) if r[1] == '' and r[0] and r[0].isupper()
    else ('' if r[0] == '' else '<tr><td>%s</td><td class="r">%s</td></tr>' % (r[0], money(r[1]) if isinstance(r[1], (int, float)) else ''))
    for r in inc_rows)
html = ('<!doctype html><meta charset=utf-8>%s<h1>AL-QAVI TRADERS</h1><h2>Income Statement — %s</h2>'
        '<table><tr><th>Head</th><th class=r>Amount</th></tr>%s'
        '<tr class="tot"><td>NET PROFIT / (LOSS)</td><td class="r">%s</td></tr></table>'
        % (css, sub, inc_html, money(rev_tot - exp_tot)))
open(os.path.join(OUT_DIR, 'Income_Statement.html'), 'w', encoding='utf-8').write(html)

rec_html = ''.join('<tr><td>%s</td><td>%s</td><td>%s</td><td class="r">%s</td></tr>'
                   % (r[0], r[2], r[3], money(r[4])) for r in rec_rows)
html = ('<!doctype html><meta charset=utf-8>%s<h1>AL-QAVI TRADERS</h1><h2>Recovery / Outstanding Receivables — %s</h2>'
        '<table><tr><th>Area</th><th>Customer</th><th>Contact</th><th class=r>Balance Due</th></tr>%s'
        '<tr class="tot"><td colspan=3>TOTAL OUTSTANDING</td><td class="r">%s</td></tr></table>'
        % (css, sub, rec_html, money(rec_total)))
open(os.path.join(OUT_DIR, 'Recovery.html'), 'w', encoding='utf-8').write(html)

# Printable index
idx = ('<!doctype html><meta charset=utf-8>%s<h1>AL-QAVI TRADERS</h1>'
       '<h2>Reports — %s</h2><table><tr><th>Report</th><th>File</th></tr>'
       '<tr><td>Income Statement (printable)</td><td><a href="Income_Statement.html">Income_Statement.html</a></td></tr>'
       '<tr><td>Recovery / Receivables (printable)</td><td><a href="Recovery.html">Recovery.html</a></td></tr>'
       '<tr><td>All reports (Excel)</td><td>AlQavi_Reports.xlsx</td></tr>'
       '<tr><td>Customer Ledgers (Excel)</td><td>AlQavi_Customer_Ledgers.xlsx</td></tr>'
       '</table>' % (css, sub))
open(os.path.join(OUT_DIR, 'index.html'), 'w', encoding='utf-8').write(idx)

# ── Console summary ────────────────────────────────────────────────────
print('OUTPUT FOLDER :', OUT_DIR)
print('  AlQavi_Reports.xlsx        (', len(wb.sheetnames), 'sheets )')
print('  AlQavi_Customer_Ledgers.xlsx (', len(ar_parties), 'party ledgers )')
print('  index.html, Income_Statement.html, Recovery.html (printable)')
print('-' * 56)
print('Accounts in ledger      :', len(bal))
print('Sales invoices          :', len(sa))
print('Total Net Sales         : {:>16,.2f}'.format(s_amt))
print('Total Revenue           : {:>16,.2f}'.format(rev_tot))
print('Total Expenses          : {:>16,.2f}'.format(exp_tot))
print('NET PROFIT / (LOSS)     : {:>16,.2f}'.format(rev_tot - exp_tot))
print('Outstanding Receivables : {:>16,.2f}   ({} customers)'.format(rec_total, len(rec_rows)))
print('Purchases (net)         : {:>16,.2f}   ({} bills)'.format(p_net, len(pr_rows)))
print('Stock value (at cost)   : {:>16,.2f}'.format(st_cost))
print('Total Assets            : {:>16,.2f}'.format(assets))
